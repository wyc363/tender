from __future__ import annotations

import argparse
import json
import random
import re
from collections import Counter
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

NUMBER_RE = re.compile(r"(?m)^\s*(\d+(?:\.\d+)+)(?=\s|[\u4e00-\u9fff（(])")
BAD_UNIT_RE = re.compile(r"^(?:\s*){no}\s*(?:km|mm|cm|m²|m2|m|㎡|%|元|万元|亿元|天|年|月)(?:\b|。|，|,|：|:)", re.I)
DATE_NO_RE = re.compile(r"^(?:19|20)\d{2}\.(?:0?[1-9]|1[0-2])$")
LOW_CONFIDENCE = 0.70

HEADERS = ["审核状态","人工问题类型","人工备注","doc_id","chunk_id","chunk_type","section_path","clause_no","clause_title",
           "pdf_page_start","pdf_page_end","跨页页数","text_raw","text_length","text_for_retrieval","structure_confidence","quality_flags",
           "source_block_ids","source_spans","检出的编号","编号数量","疑似多Clause","疑似错误Clause编号","疑似过短","疑似超长",
           "是否UNCERTAIN","自动审核提示"]
STATUS_VALUES = '"未审核,正确,有问题,无法确定"'
ISSUE_VALUES = '"Clause边界错误,Clause编号错误,Table Row错误,正文/表格类型错误,跨页错误,正文缺失,内容重复,页码错误,UNCERTAIN误判,其他"'

def load_jsonl(path: Path) -> list[dict]:
    return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]

def readable(value) -> str:
    if value is None: return ""
    if isinstance(value, list): return " > ".join(map(str,value))
    return str(value)

def analyze(chunk: dict) -> dict:
    text=chunk.get("text_raw") or ""; numbers=NUMBER_RE.findall(text); clause_no=str(chunk.get("clause_no") or "")
    depth=clause_no.count(".") if clause_no else None
    same_level=[n for n in numbers if depth is not None and n.count(".")==depth]
    multi=chunk.get("chunk_type")=="CLAUSE" and len(dict.fromkeys(same_level))>=2
    bad=False
    if clause_no:
        bad=bool(re.search(BAD_UNIT_RE.pattern.format(no=re.escape(clause_no)),text,re.I) or DATE_NO_RE.fullmatch(clause_no))
    length=len(text); short=length<20; long=length>1800; uncertain=chunk.get("chunk_type")=="UNCERTAIN_FRAGMENT"
    cross=max(1,int(chunk.get("pdf_page_end",0))-int(chunk.get("pdf_page_start",0))+1)
    flags=chunk.get("quality_flags") or []; conf=chunk.get("structure_confidence")
    risks=[]; hints=[]
    if multi: risks.append("多Clause"); hints.append("一个 CLAUSE Chunk 中检测到多个疑似同级条款编号，请检查是否发生条款合并。")
    if bad: risks.append("错误Clause编号"); hints.append("clause_no 后紧跟计量单位/百分号/日期等，可能把数值误识别为条款号。")
    if short: risks.append("过短"); hints.append("Chunk 很短，请确认是否仍构成完整独立业务要求。")
    if long: risks.append("超长"); hints.append("超过当前 hard_max_chars，请检查是否未按子项拆分。")
    if uncertain: risks.append("UNCERTAIN"); hints.append("UNCERTAIN Chunk，请判断正文/表格类型。"+(" quality_flags: "+", ".join(flags) if flags else ""))
    if cross>1: risks.append("跨页"); hints.append("跨页 Chunk，请对照连续 PDF 页检查是否为真实连续条款，是否存在错合。")
    if flags: risks.append("质量标记")
    if conf is not None and float(conf)<LOW_CONFIDENCE: risks.append("低置信度"); hints.append(f"structure_confidence={float(conf):.2f}，请重点核对结构类型和边界。")
    return {"numbers":numbers,"multi":multi,"bad":bad,"short":short,"long":long,"uncertain":uncertain,"cross":cross,
            "risks":list(dict.fromkeys(risks)),"hint":" ".join(dict.fromkeys(hints)),"length":length}

def row_for(chunk: dict, a: dict) -> list:
    return ["未审核","","",chunk.get("doc_id"),chunk.get("chunk_id"),chunk.get("chunk_type"),readable(chunk.get("section_path")),
            chunk.get("clause_no") or "",chunk.get("clause_title") or "",chunk.get("pdf_page_start"),chunk.get("pdf_page_end"),a["cross"],
            chunk.get("text_raw") or "",a["length"],chunk.get("text_for_retrieval") or "",chunk.get("structure_confidence"),
            ", ".join(chunk.get("quality_flags") or []),", ".join(chunk.get("source_block_ids") or []),
            json.dumps(chunk.get("source_spans") or [],ensure_ascii=False,separators=(",",":")),"; ".join(a["numbers"]),len(a["numbers"]),
            "是" if a["multi"] else "否","是" if a["bad"] else "否","是" if a["short"] else "否","是" if a["long"] else "否",
            "是" if a["uncertain"] else "否",a["hint"]]

def style_table(ws, widths: dict[int,float], risk=False):
    ws.freeze_panes="A2"; ws.auto_filter.ref=ws.dimensions; ws.sheet_view.showGridLines=False
    fill=PatternFill("solid",fgColor="1F4E78"); font=Font(color="FFFFFF",bold=True)
    for c in ws[1]: c.fill=fill; c.font=font; c.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
    compact=ws.title=="统计汇总"
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment=Alignment(vertical="top",wrap_text=True)
        ws.row_dimensions[row[0].row].height=24 if compact else 60
    for col,width in widths.items(): ws.column_dimensions[ws.cell(1,col).column_letter].width=width
    if ws.max_row>1:
        header_values=[c.value for c in ws[1]]
        status_col=header_values.index("审核状态")+1 if "审核状态" in header_values else (header_values.index("人工判断")+1 if "人工判断" in header_values else None)
        if status_col:
            letter=ws.cell(1,status_col).column_letter; dv=DataValidation(type="list",formula1=STATUS_VALUES,allow_blank=False); ws.add_data_validation(dv); dv.add(f"{letter}2:{letter}{ws.max_row}")
        issue_col=header_values.index("人工问题类型")+1 if "人工问题类型" in header_values else None
        if issue_col:
            letter=ws.cell(1,issue_col).column_letter; dv2=DataValidation(type="list",formula1=ISSUE_VALUES,allow_blank=True); ws.add_data_validation(dv2); dv2.add(f"{letter}2:{letter}{ws.max_row}")
        if risk: ws.conditional_formatting.add(f"A2:{ws.cell(ws.max_row,ws.max_column).coordinate}",FormulaRule(formula=["$A2<>\"\""],fill=PatternFill("solid",fgColor="FFF2CC")))

def spread_sample(items: list[dict], n: int, rng: random.Random) -> list[dict]:
    if len(items)<=n:return list(items)
    ordered=sorted(items,key=lambda x:(x["chunk"].get("pdf_page_start",0),x["chunk"].get("chunk_id","")))
    thirds=[ordered[:len(ordered)//3],ordered[len(ordered)//3:2*len(ordered)//3],ordered[2*len(ordered)//3:]]
    counts=[n//3,n//3,n-2*(n//3)]; out=[]
    for group,k in zip(thirds,counts): out.extend(rng.sample(group,min(k,len(group))))
    return out

def main() -> int:
    p=argparse.ArgumentParser(); p.add_argument("--stage2-root",type=Path,required=True); p.add_argument("--docs",nargs="+",required=True); p.add_argument("--output",type=Path,required=True); a=p.parse_args()
    records={}; blocks={}; summary={}
    for doc in a.docs:
        chunks=load_jsonl(a.stage2_root/doc/"chunks.jsonl"); records[doc]=[{"chunk":c,"analysis":analyze(c)} for c in chunks]
        blocks[doc]={b["block_id"]:b for b in load_jsonl(a.stage2_root/doc/"region_blocks.jsonl")}
        ct=Counter(c["chunk_type"] for c in chunks); aa=[x["analysis"] for x in records[doc]]
        summary[doc]={"total_chunks":len(chunks),"clause":ct["CLAUSE"],"table_row":ct["TABLE_ROW"],"paragraph":ct["PARAGRAPH"],"uncertain":ct["UNCERTAIN_FRAGMENT"],
                      "suspected_multi_clause":sum(x["multi"] for x in aa),"suspected_bad_clause_no":sum(x["bad"] for x in aa),"short_chunks":sum(x["short"] for x in aa),
                      "long_chunks":sum(x["long"] for x in aa),"cross_page_chunks":sum(x["cross"]>1 for x in aa),"with_quality_flags":sum(bool(r["chunk"].get("quality_flags")) for r in records[doc])}
    wb=Workbook(); wb.remove(wb.active)
    widths={1:12,2:20,3:24,4:10,5:22,6:20,7:28,8:14,9:28,10:12,11:12,12:10,13:80,14:12,15:60,16:16,17:30,18:35,19:55,20:28,21:10,22:13,23:18,24:12,25:12,26:14,27:65}
    for doc in a.docs:
        ws=wb.create_sheet(f"{doc}_All"); ws.append(HEADERS)
        for r in records[doc]: ws.append(row_for(r["chunk"],r["analysis"]))
        style_table(ws,widths)
    focus=wb.create_sheet("重点审核"); focus.append(["风险类型"]+HEADERS)
    for doc in a.docs:
        for r in records[doc]:
            if r["analysis"]["risks"]: focus.append([" | ".join(r["analysis"]["risks"])]+row_for(r["chunk"],r["analysis"]))
    style_table(focus,{1:30,**{k+1:v for k,v in widths.items()}},True)
    unc=wb.create_sheet("UNCERTAIN专项"); uh=["doc_id","chunk_id","PDF页码","text_raw","text_length","quality_flags","structure_confidence","uncertain_reason","人工判断","人工备注"]; unc.append(uh)
    for doc in a.docs:
        for r in records[doc]:
            if not r["analysis"]["uncertain"]:continue
            c=r["chunk"]; related=[blocks[doc].get(x,{}) for x in c.get("source_block_ids",[])]; reasons=[]
            for b in related: reasons += b.get("region_evidence",[])+b.get("quality_flags",[])
            reasons=[x for x in dict.fromkeys(reasons) if x in {"PYMUPDF_TABLE_OVERLAP","NATIVE_TABLE_ONLY_UNCERTAIN"}]
            unc.append([doc,c["chunk_id"],f"{c['pdf_page_start']}-{c['pdf_page_end']}",c.get("text_raw",""),r["analysis"]["length"],", ".join(c.get("quality_flags") or []),c.get("structure_confidence"),", ".join(reasons),"未审核",""])
    style_table(unc,{1:10,2:22,3:12,4:90,5:12,6:30,7:16,8:34,9:14,10:30},True)
    stat=wb.create_sheet("统计汇总"); stat.append(["指标"]+a.docs)
    labels=[("Chunk 总数","total_chunks"),("CLAUSE 数","clause"),("TABLE_ROW 数","table_row"),("PARAGRAPH 数","paragraph"),("UNCERTAIN_FRAGMENT 数","uncertain"),("CROSS_PAGE Chunk 数","cross_page_chunks"),("疑似多条款 Chunk 数","suspected_multi_clause"),("疑似错误 clause_no 数","suspected_bad_clause_no"),("超长 Chunk 数","long_chunks"),("极短 Chunk 数","short_chunks"),("有 quality_flags 的 Chunk 数","with_quality_flags")]
    for label,key in labels: stat.append([label]+[summary[d][key] for d in a.docs])
    style_table(stat,{1:32,2:16,3:16}); stat.freeze_panes=None; stat.auto_filter.ref=stat.dimensions
    sample=wb.create_sheet("分层抽查"); sample.append(["抽样类别"]+HEADERS); rng=random.Random(42)
    for doc in a.docs:
        chosen=set()
        groups=[("CLAUSE",lambda r:r["chunk"]["chunk_type"]=="CLAUSE",10),("TABLE_ROW",lambda r:r["chunk"]["chunk_type"]=="TABLE_ROW",10),("CROSS_PAGE",lambda r:r["analysis"]["cross"]>1,5),("UNCERTAIN",lambda r:r["analysis"]["uncertain"],10),("SHORT",lambda r:r["analysis"]["short"],5),("LONG",lambda r:r["analysis"]["long"],5)]
        for label,pred,n in groups:
            pool=[r for r in records[doc] if pred(r) and r["chunk"]["chunk_id"] not in chosen]
            picks=spread_sample(pool,n,rng) if label=="UNCERTAIN" else rng.sample(pool,min(n,len(pool)))
            for r in picks: chosen.add(r["chunk"]["chunk_id"]); sample.append([f"{doc}-{label}"]+row_for(r["chunk"],r["analysis"]))
        if doc=="DOC011":
            pool=[r for r in records[doc] if r["analysis"]["uncertain"] and r["chunk"]["chunk_id"] not in chosen]
            for r in spread_sample(pool,30,rng): chosen.add(r["chunk"]["chunk_id"]); sample.append(["DOC011-UNCERTAIN额外"]+row_for(r["chunk"],r["analysis"]))
    style_table(sample,{1:24,**{k+1:v for k,v in widths.items()}},True)
    a.output.parent.mkdir(parents=True,exist_ok=True); wb.save(a.output)
    summary_path=a.output.with_name("stage2_step1_chunk_review_summary.json"); summary_path.write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding="utf-8")
    # Re-open to verify the produced file, sheet contract, panes, filters and row counts.
    check=load_workbook(a.output,read_only=False); expected=[f"{d}_All" for d in a.docs]+["重点审核","UNCERTAIN专项","统计汇总","分层抽查"]
    assert check.sheetnames==expected
    assert all(check[f"{d}_All"].max_row-1==summary[d]["total_chunks"] for d in a.docs)
    assert all(check[f"{d}_All"].freeze_panes=="A2" and check[f"{d}_All"].auto_filter.ref for d in a.docs)
    check.close(); print(json.dumps({"output":str(a.output),"summary":summary},ensure_ascii=False)); return 0
if __name__=="__main__": raise SystemExit(main())
