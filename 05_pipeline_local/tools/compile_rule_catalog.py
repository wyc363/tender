from __future__ import annotations

import argparse
import hashlib
import json
import sys
from collections import OrderedDict
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


SCHEMA_VERSION = "1.0"
COMPILER_VERSION = "1.0"
REQUIRED_HEADERS = [
    "规则序号", "审查类型", "审查点名称", "行业", "风险等级", "审查规则说明",
    "地域范围", "违规条目", "合法条目", "政策来源", "规则适用开始日期",
    "规则适用结束日期", "扫描范围", "备注",
]


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def rule_no(value: Any) -> str:
    if value is None or value == "":
        return ""
    if isinstance(value, (int, float)):
        return f"{float(value):.1f}"
    raw = str(value).strip()
    try:
        return f"{float(raw):.1f}"
    except ValueError:
        return raw


def unique(values: list[str]) -> list[str]:
    return list(dict.fromkeys(value for value in values if value))


def catalog_order_from_rule(rule_id: str) -> int:
    try:
        number = int(float(rule_id))
    except ValueError as exc:
        raise ValueError(f"规则序号不是数值编号: {rule_id!r}") from exc
    order = number // 100
    if order < 1:
        raise ValueError(f"无法从规则序号取得审查点顺序: {rule_id!r}")
    return order


def load_overrides(path: Path) -> dict[str, str]:
    if not path.is_file():
        return {}
    value = json.loads(path.read_text(encoding="utf-8"))
    by_name = value.get("by_name") or {}
    if not isinstance(by_name, dict):
        raise ValueError("review_point_id_overrides.json 的 by_name 必须是对象")
    return {str(k): str(v) for k, v in by_name.items()}


def build_catalog(source: Path, sheet_name: str, overrides_path: Path) -> tuple[dict[str, Any], dict[str, Any]]:
    workbook = load_workbook(source, read_only=True, data_only=True)
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"工作表不存在: {sheet_name}")
    sheet = workbook[sheet_name]
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        raise ValueError("规则工作表为空")
    headers = [text(value) for value in rows[0]]
    missing = [name for name in REQUIRED_HEADERS if name not in headers]
    if missing:
        raise ValueError(f"缺少必需列: {missing}")
    column = {name: headers.index(name) for name in REQUIRED_HEADERS}
    overrides = load_overrides(overrides_path)
    groups: OrderedDict[str, list[dict[str, Any]]] = OrderedDict()
    errors: list[str] = []
    warnings: list[str] = []
    seen_rule_ids: dict[str, int] = {}

    for excel_row, values in enumerate(rows[1:], start=2):
        values = list(values) + [None] * (len(headers) - len(values))
        rid = rule_no(values[column["规则序号"]])
        name = text(values[column["审查点名称"]])
        if not rid and not name:
            continue
        if not rid:
            errors.append(f"第{excel_row}行缺少规则序号")
        if not name:
            errors.append(f"第{excel_row}行缺少审查点名称")
        if rid in seen_rule_ids:
            errors.append(f"规则序号重复: {rid}（第{seen_rule_ids[rid]}行和第{excel_row}行）")
        seen_rule_ids[rid] = excel_row
        record = {header: text(values[index]) for header, index in column.items()}
        record["规则序号"] = rid
        record["source_row"] = excel_row
        groups.setdefault(name, []).append(record)

    points: list[dict[str, Any]] = []
    used_ids: dict[str, str] = {}
    used_orders: dict[int, str] = {}
    for name, records in groups.items():
        orders = unique([str(catalog_order_from_rule(row["规则序号"])) for row in records if row["规则序号"]])
        if len(orders) != 1:
            errors.append(f"审查点{name!r}包含多个目录序号: {orders}")
            continue
        order = int(orders[0])
        if order in used_orders and used_orders[order] != name:
            errors.append(f"目录序号{order}同时对应{name!r}和{used_orders[order]!r}")
        used_orders[order] = name
        review_point_id = overrides.get(name, f"RP_{order:03d}")
        if review_point_id in used_ids and used_ids[review_point_id] != name:
            errors.append(f"review_point_id重复: {review_point_id}")
        used_ids[review_point_id] = name

        constants: dict[str, str] = {}
        for header, output_name in (
            ("审查类型", "type"), ("风险等级", "risk_level"), ("审查规则说明", "rule")
        ):
            values = unique([row[header] for row in records])
            if len(values) != 1:
                errors.append(f"审查点{name!r}的{header}不唯一: {values}")
            constants[output_name] = values[0] if values else ""
            if not constants[output_name]:
                errors.append(f"审查点{name!r}缺少{header}")

        entries = []
        for row in records:
            if not row["违规条目"]:
                errors.append(f"第{row['source_row']}行违规条目为空")
            entries.append({
                "rule_no": row["规则序号"],
                "violation_term": row["违规条目"],
                "legal_boundary": row["合法条目"],
                "industry": row["行业"],
                "region": row["地域范围"],
                "policy": row["政策来源"],
                "effective_start": row["规则适用开始日期"] or None,
                "effective_end": row["规则适用结束日期"] or None,
                "scan_scope_hint": row["扫描范围"],
                "remark": row["备注"],
                "source_row": row["source_row"],
            })

        rule_ids = [entry["rule_no"] for entry in entries]
        violation_terms = [entry["violation_term"] for entry in entries]
        if len(rule_ids) != len(violation_terms):
            errors.append(f"审查点{name!r}的规则编号与违规条目无法一一对应")
        points.append({
            "review_point_id": review_point_id,
            "catalog_order": order,
            "index": order,
            "name": name,
            **constants,
            "rule_ids": rule_ids,
            "rule_count": len(entries),
            "industry": unique([entry["industry"] for entry in entries]),
            "region": unique([entry["region"] for entry in entries]),
            "violation_terms": violation_terms,
            "legal_terms": unique([entry["legal_boundary"] for entry in entries]),
            "policy_sources": unique([entry["policy"] for entry in entries]),
            "scan_scopes": unique([entry["scan_scope_hint"] for entry in entries]),
            "effective_start_dates": unique([entry["effective_start"] or "" for entry in entries]),
            "effective_end_dates": unique([entry["effective_end"] or "" for entry in entries]),
            "source_rows": [entry["source_row"] for entry in entries],
            "rule_entries": entries,
        })

    points.sort(key=lambda row: row["catalog_order"])
    orders = [row["catalog_order"] for row in points]
    expected_orders = list(range(1, len(points) + 1))
    if orders != expected_orders:
        warnings.append(f"审查点目录序号不是连续的1..N: {orders}")

    catalog = {
        "status": "AUTHORITATIVE_RULE_CATALOG",
        "schema_version": SCHEMA_VERSION,
        "compiler_version": COMPILER_VERSION,
        "source": {
            "file_name": source.name,
            "source_sha256": sha256_file(source),
            "sheet": sheet_name,
            "data_rows": sum(len(group) for group in groups.values()),
        },
        # Legacy metadata is retained for existing callers.
        "source_file": source.name,
        "source_sha256": sha256_file(source),
        "sheet": sheet_name,
        "leaf_rule_count": sum(len(group) for group in groups.values()),
        "review_type_count": len({row["type"] for row in points}),
        "review_point_count": len(points),
        "review_points": points,
    }
    report = {
        "status": "PASS" if not errors else "FAIL",
        "schema_version": SCHEMA_VERSION,
        "source_file": str(source),
        "source_sha256": catalog["source_sha256"],
        "sheet": sheet_name,
        "data_rows": catalog["leaf_rule_count"],
        "review_points": catalog["review_point_count"],
        "review_types": catalog["review_type_count"],
        "rule_ids_unique": len(seen_rule_ids) == catalog["leaf_rule_count"],
        "review_point_ids_unique": len(used_ids) == catalog["review_point_count"],
        "errors": errors,
        "warnings": warnings,
    }
    return catalog, report


def write_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def write_preview(path: Path, catalog: dict[str, Any]) -> None:
    lines = [
        "# 招标文件审查规则目录",
        "",
        f"- 来源文件：`{catalog['source_file']}`",
        f"- 来源 SHA256：`{catalog['source_sha256']}`",
        f"- 审查类型：{catalog['review_type_count']}",
        f"- 审查点：{catalog['review_point_count']}",
        f"- 叶子规则：{catalog['leaf_rule_count']}",
        "",
    ]
    for point in catalog["review_points"]:
        lines.extend([
            f"## {point['catalog_order']}. {point['name']}", "",
            f"- 内部 ID：`{point['review_point_id']}`",
            f"- 审查类型：{point['type']}",
            f"- 风险等级：{point['risk_level']}",
            f"- 审查规则：{point['rule']}",
            f"- 叶子规则数：{point['rule_count']}", "",
            "### 违规条目", "",
        ])
        for entry in point["rule_entries"]:
            term = entry["violation_term"].replace("\n", "<br>")
            lines.append(f"- `{entry['rule_no']}`（Excel第{entry['source_row']}行）：{term}")
        lines.append("")
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines), encoding="utf-8")


def main() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    project = Path(__file__).resolve().parents[1]
    parser = argparse.ArgumentParser(description="Compile the authoritative Excel rule workbook into runtime JSON")
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--sheet", default="审查规则")
    parser.add_argument("--output", type=Path, default=project / "config" / "rule_catalog.compiled.json")
    parser.add_argument("--preview", type=Path, default=project / "config" / "rule_catalog.preview.md")
    parser.add_argument("--validation", type=Path, default=project / "config" / "rule_catalog.validation.json")
    parser.add_argument("--id-overrides", type=Path, default=project / "config" / "review_point_id_overrides.json")
    args = parser.parse_args()
    source = args.input.expanduser().resolve()
    if not source.is_file():
        raise SystemExit(f"Excel rule workbook not found: {source}")
    try:
        catalog, report = build_catalog(source, args.sheet, args.id_overrides.resolve())
        report["generated_at_utc"] = datetime.now(timezone.utc).isoformat()
        write_json(args.validation.resolve(), report)
        if report["status"] != "PASS":
            print(json.dumps(report, ensure_ascii=False, indent=2))
            raise SystemExit(1)
        write_json(args.output.resolve(), catalog)
        write_preview(args.preview.resolve(), catalog)
        print(json.dumps({
            "status": "PASS",
            "catalog": str(args.output.resolve()),
            "preview": str(args.preview.resolve()),
            "validation": str(args.validation.resolve()),
            "review_points": report["review_points"],
            "leaf_rules": report["data_rows"],
            "source_sha256": report["source_sha256"],
        }, ensure_ascii=False, indent=2))
    except Exception as exc:
        failure = {
            "status": "FAIL",
            "generated_at_utc": datetime.now(timezone.utc).isoformat(),
            "source_file": str(source),
            "errors": [str(exc)],
        }
        write_json(args.validation.resolve(), failure)
        print(json.dumps(failure, ensure_ascii=False, indent=2))
        raise SystemExit(1)


if __name__ == "__main__":
    main()
