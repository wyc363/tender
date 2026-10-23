from __future__ import annotations

import hashlib
import json
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


CORE_RULES = {6, 7, 33, 34, 35, 36, 37, 38, 39, 40}
PARTIAL_RULES = {8, 9, 19, 41}
CROSS_REVIEW_POINT_RULES = {46, 47}


def file_sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest().upper()


def _value(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return value


def _tier(rule_no: int) -> str | None:
    if rule_no in CORE_RULES:
        return "CORE"
    if rule_no in PARTIAL_RULES:
        return "PARTIAL_RELATED"
    if rule_no in CROSS_REVIEW_POINT_RULES:
        return "CROSS_REVIEW_POINT"
    return None


def _mapping_rows(mapping_workbook: Path) -> dict[int, list[dict[str, Any]]]:
    wb = load_workbook(mapping_workbook, data_only=True, read_only=False)
    ws = wb["旧新映射候选"]
    headers = [str(c.value or "").strip() for c in ws[1]]
    result: dict[int, list[dict[str, Any]]] = {}
    for values in ws.iter_rows(min_row=2, values_only=True):
        row = dict(zip(headers, values))
        try:
            new_id = int(row.get("新版规则序号"))
        except (TypeError, ValueError):
            continue
        result.setdefault(new_id, []).append(
            {
                "legacy_rule_id": row.get("旧RuleID"),
                "mapping_type": row.get("映射类型（初步表）"),
                "semantic_note": row.get("业务语义对照说明"),
                "approval_status": row.get("映射批准状态"),
                "production_allowed": bool(row.get("允许生产切换")),
            }
        )
    return result


def build_scope_registry(new_rules_workbook: Path, mapping_workbook: Path) -> dict[str, Any]:
    mappings = _mapping_rows(mapping_workbook)
    wb = load_workbook(new_rules_workbook, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    headers = [str(x or "").strip() for x in next(ws.iter_rows(values_only=True))]
    rules: list[dict[str, Any]] = []
    for row_no, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        row = dict(zip(headers, values))
        try:
            rule_no = int(row.get("规则序号"))
        except (TypeError, ValueError):
            continue
        tier = _tier(rule_no)
        if not tier:
            continue
        rules.append(
            {
                "rule_id": f"20261013:{rule_no}",
                "source_rule_number": rule_no,
                "scope_tier": tier,
                "review_type": row.get("审查类型"),
                "review_point": row.get("审查点"),
                "industry": row.get("行业"),
                "risk_level": row.get("风险等级"),
                "rule_text": row.get("审查规则说明"),
                "scan_scope": row.get("扫描范围"),
                "violation_examples": row.get("违规条目"),
                "legal_examples": row.get("合法条目"),
                "policy_source": row.get("政策来源"),
                "source_start_date": _value(row.get("政策适用开始日期")),
                "source_end_date": _value(row.get("政策适用结束日期")),
                "source_sheet": ws.title,
                "source_excel_row": row_no,
                "source_file_sha256": file_sha256(new_rules_workbook),
                "policy_verified": False,
                "policy_verification_status": "SOURCE_TABLE_ONLY_NOT_OFFICIAL_POLICY_VERIFIED",
                "production_authorized": False,
                "legacy_mappings": mappings.get(rule_no, []),
                "scope_note": (
                    "RP12主干候选" if tier == "CORE" else
                    "与旧RP12部分相关，需业务确认后才能进入评价范围" if tier == "PARTIAL_RELATED" else
                    "邻近审查项风险，仅记录cross_rule_links，不进入RP12结论"
                ),
            }
        )
    return {
        "registry_id": "rp12-20261013-scope-registry-v0.1-shadow",
        "status": "PROVISIONAL_SHADOW_ONLY",
        "review_point_id": "RP_REGION_INDUSTRY_AWARD",
        "client_index": 12,
        "source_workbook": str(new_rules_workbook),
        "source_workbook_sha256": file_sha256(new_rules_workbook),
        "mapping_workbook": str(mapping_workbook),
        "mapping_workbook_sha256": file_sha256(mapping_workbook),
        "rules": rules,
        "counts": {
            "core": sum(r["scope_tier"] == "CORE" for r in rules),
            "partial_related": sum(r["scope_tier"] == "PARTIAL_RELATED" for r in rules),
            "cross_review_point": sum(r["scope_tier"] == "CROSS_REVIEW_POINT" for r in rules),
        },
        "warnings": [
            "policy_verified=false表示只核对了规则Excel，未独立核验官方政策原文和现行效力。",
            "规则34原表措辞数据异常应原样保留并交规则负责人核对。",
            "35-40不得在非福建水利项目中直接适用。",
        ],
    }


def write_registry(registry: dict[str, Any], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(registry, ensure_ascii=False, indent=2), encoding="utf-8")
